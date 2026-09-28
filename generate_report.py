import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_NAME = "iam_database.db"
OUTPUT_FILE = "access_review_report.xlsx"

def generate_excel_report():
    conn = sqlite3.connect(DB_NAME)

    # 1. Summary Sheet
    summary_data = [
        ["Metric", "Count / Value", "Risk Level", "Description"],
        ["Total Employees", pd.read_sql_query("SELECT COUNT(*) FROM employees", conn).iloc[0, 0], "Info", "Total employee headcount"],
        ["Active Employees", pd.read_sql_query("SELECT COUNT(*) FROM employees WHERE status='active'", conn).iloc[0, 0], "Info", "Currently employed staff"],
        ["Leavers with Active Access", pd.read_sql_query("SELECT COUNT(DISTINCT aa.employee_id) FROM access_assignments aa JOIN employees e ON aa.employee_id=e.employee_id WHERE e.status='leaver'", conn).iloc[0, 0], "HIGH", "Former employees holding access"],
        ["Dormant Access (180+ Days)", pd.read_sql_query("SELECT COUNT(*) FROM access_assignments WHERE last_used_date IS NOT NULL AND JULIANDAY('2026-09-01') - JULIANDAY(last_used_date) >= 180", conn).iloc[0, 0], "MEDIUM", "Access accounts unused in 6+ months"],
        ["Unused Applications", pd.read_sql_query("SELECT COUNT(*) FROM (SELECT a.app_id FROM applications a LEFT JOIN entitlements ent ON a.app_id=ent.app_id LEFT JOIN access_assignments aa ON ent.entitlement_id=aa.entitlement_id LEFT JOIN employees e ON aa.employee_id=e.employee_id AND e.status='active' GROUP BY a.app_id HAVING COUNT(e.employee_id)=0)", conn).iloc[0, 0], "LOW", "Applications with zero active users"],
        ["Mover Privilege Creep Issues", pd.read_sql_query("SELECT COUNT(*) FROM access_assignments aa JOIN employees e ON aa.employee_id=e.employee_id JOIN entitlements ent ON aa.entitlement_id=ent.entitlement_id JOIN applications a ON ent.app_id=a.app_id WHERE e.status='mover' AND a.owner_department != e.department AND a.owner_department != 'IT Ops'", conn).iloc[0, 0], "MEDIUM", "Movers retaining old department access"]
    ]
    df_summary = pd.DataFrame(summary_data[1:], columns=summary_data[0])

    # 2. Leavers with Access Sheet
    q_leavers = """
    SELECT e.employee_id, e.name, e.department, e.leave_date, a.app_name, ent.entitlement_name, aa.last_used_date
    FROM access_assignments aa
    JOIN employees e ON aa.employee_id = e.employee_id
    JOIN entitlements ent ON aa.entitlement_id = ent.entitlement_id
    JOIN applications a ON ent.app_id = a.app_id
    WHERE e.status = 'leaver'
    ORDER BY e.leave_date DESC;
    """
    df_leavers = pd.read_sql_query(q_leavers, conn)

    # 3. Unused Access Sheet
    q_unused = """
    SELECT e.employee_id, e.name, e.department, a.app_name, ent.entitlement_name, aa.last_used_date,
           CAST((JULIANDAY('2026-09-01') - JULIANDAY(aa.last_used_date)) AS INTEGER) AS days_inactive
    FROM access_assignments aa
    JOIN employees e ON aa.employee_id = e.employee_id
    JOIN entitlements ent ON aa.entitlement_id = ent.entitlement_id
    JOIN applications a ON ent.app_id = a.app_id
    WHERE aa.last_used_date IS NOT NULL AND JULIANDAY('2026-09-01') - JULIANDAY(aa.last_used_date) >= 180
    ORDER BY days_inactive DESC;
    """
    df_unused = pd.read_sql_query(q_unused, conn)

    # 4. Cleanup Candidates Sheet
    q_cleanup = """
    SELECT a.app_id, a.app_name, a.owner_department, a.status AS app_status
    FROM applications a
    LEFT JOIN entitlements ent ON a.app_id = ent.app_id
    LEFT JOIN access_assignments aa ON ent.entitlement_id = aa.entitlement_id
    LEFT JOIN employees e ON aa.employee_id = e.employee_id AND e.status = 'active'
    GROUP BY a.app_id, a.app_name, a.owner_department, a.status
    HAVING COUNT(e.employee_id) = 0;
    """
    df_cleanup = pd.read_sql_query(q_cleanup, conn)

    # 5. Mover Issues Sheet
    q_movers = """
    SELECT e.employee_id, e.name, e.department AS current_department, a.owner_department AS app_owner_department,
           a.app_name, ent.entitlement_name, aa.granted_date
    FROM access_assignments aa
    JOIN employees e ON aa.employee_id = e.employee_id
    JOIN entitlements ent ON aa.entitlement_id = ent.entitlement_id
    JOIN applications a ON ent.app_id = a.app_id
    WHERE e.status = 'mover' AND a.owner_department != e.department AND a.owner_department != 'IT Ops';
    """
    df_movers = pd.read_sql_query(q_movers, conn)

    conn.close()

    # 6. Proposed Business Roles Sheet
    from role_mining import mine_business_roles
    df_roles = mine_business_roles()

    # --- Write to Excel using pandas ExcelWriter ---
    with pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl') as writer:
        df_summary.to_excel(writer, sheet_name='Summary', index=False)
        df_leavers.to_excel(writer, sheet_name='Leavers with Access', index=False)
        df_unused.to_excel(writer, sheet_name='Unused Access', index=False)
        df_cleanup.to_excel(writer, sheet_name='Cleanup Candidates', index=False)
        df_movers.to_excel(writer, sheet_name='Mover Issues', index=False)
        df_roles.to_excel(writer, sheet_name='Proposed Business Roles', index=False)

    # --- openpyxl Formatting ---
    wb = openpyxl.load_workbook(OUTPUT_FILE)
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    high_risk_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        ws.freeze_panes = "A2"
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                cell.border = thin_border
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        if sheetname == "Leavers with Access":
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.fill = high_risk_fill

    wb.save(OUTPUT_FILE)
    print("SUCCESS: access_review_report.xlsx has been created!")

if __name__ == "__main__":
    generate_excel_report()